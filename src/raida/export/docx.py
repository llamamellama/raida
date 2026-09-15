"""Markdown to .docx with python-docx: headings, paragraphs with inline emphasis, lists,
quotes, code blocks, tables and rules."""

from __future__ import annotations

import io
from typing import Any

from raida.export.markdown import parser


def _add_inline(paragraph: Any, children: list[Any]) -> None:
    bold = italic = code = False
    for tok in children or []:
        t = tok.type
        if t == "text":
            run = paragraph.add_run(tok.content)
            run.bold, run.italic = bold, italic
            if code:
                run.font.name = "Menlo"
        elif t == "code_inline":
            run = paragraph.add_run(tok.content)
            run.font.name = "Menlo"
        elif t == "softbreak":
            paragraph.add_run(" ")
        elif t == "hardbreak":
            paragraph.add_run().add_break()
        elif t in ("strong_open", "strong_close"):
            bold = t.endswith("open")
        elif t in ("em_open", "em_close"):
            italic = t.endswith("open")
        elif t == "link_open":
            pass
        elif t == "image":
            paragraph.add_run(f"[image: {tok.attrGet('alt') or tok.attrGet('src') or ''}]")


def md_to_docx(markdown: str, title: str) -> bytes:
    import docx
    from docx.shared import Pt

    document = docx.Document()
    document.core_properties.title = title
    tokens = parser().parse(markdown)
    list_stack: list[str] = []
    in_quote = False
    table_rows: list[list[list[Any]]] | None = None
    current_row: list[list[Any]] | None = None

    i = 0
    while i < len(tokens):
        tok = tokens[i]
        t = tok.type
        if t == "heading_open":
            level = int(tok.tag[1])
            inline = tokens[i + 1]
            p = document.add_heading(level=min(level, 6))
            _add_inline(p, inline.children or [])
            i += 3
            continue
        if t in ("bullet_list_open", "ordered_list_open"):
            list_stack.append("List Bullet" if t.startswith("bullet") else "List Number")
        elif t in ("bullet_list_close", "ordered_list_close"):
            list_stack.pop()
        elif t == "blockquote_open":
            in_quote = True
        elif t == "blockquote_close":
            in_quote = False
        elif t == "paragraph_open":
            inline = tokens[i + 1]
            if list_stack:
                style = list_stack[-1] + (f" {len(list_stack)}" if len(list_stack) > 1 else "")
                try:
                    p = document.add_paragraph(style=style)
                except KeyError:
                    p = document.add_paragraph(style=list_stack[-1])
            elif in_quote:
                p = document.add_paragraph(style="Intense Quote")
            else:
                p = document.add_paragraph()
            _add_inline(p, inline.children or [])
            i += 3
            continue
        elif t in ("fence", "code_block"):
            p = document.add_paragraph()
            run = p.add_run(tok.content.rstrip("\n"))
            run.font.name = "Menlo"
            run.font.size = Pt(9)
        elif t == "hr":
            document.add_paragraph("_" * 40)
        elif t == "table_open":
            table_rows = []
        elif t == "tr_open":
            current_row = []
        elif t in ("th_open", "td_open"):
            inline = tokens[i + 1]
            if current_row is not None:
                current_row.append(inline.children or [])
            i += 3
            continue
        elif t == "tr_close":
            if table_rows is not None and current_row is not None:
                table_rows.append(current_row)
            current_row = None
        elif t == "table_close" and table_rows:
            cols = max(len(r) for r in table_rows)
            table = document.add_table(rows=len(table_rows), cols=cols)
            table.style = "Table Grid"
            for r, row in enumerate(table_rows):
                for c, cell_children in enumerate(row):
                    cell = table.cell(r, c)
                    cell.paragraphs[0].text = ""
                    _add_inline(cell.paragraphs[0], cell_children)
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            table_rows = None
        i += 1

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
